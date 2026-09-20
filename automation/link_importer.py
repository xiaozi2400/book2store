"""
Excel链接导入器 - 从百度网盘Excel导入分享链接
"""

import os
import sys
from pathlib import Path
from typing import Dict, Optional

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from .config import config
from .database import DatabaseManager
from .utils import ensure_dir, logger, match_filename_to_book


class LinkImporter:
    """Excel链接导入器"""

    def __init__(self):
        self.db = DatabaseManager()
        self.output_dir = Path(config.output_dir)

    def import_from_excel(self, excel_path: str) -> Dict:
        """从Excel导入分享链接"""
        logger.info(f"开始导入Excel: {excel_path}")

        try:
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]
            logger.info(f"Excel表头: {headers}")

            results = {"total": 0, "matched": 0, "unmatched": 0, "details": []}

            for row_idx in range(2, ws.max_row + 1):
                row_data = {}
                for col_idx, header in enumerate(headers, 1):
                    cell_value = ws.cell(row_idx, col_idx).value
                    row_data[header] = cell_value

                results["total"] += 1

                matched_book = self._match_book(row_data)

                if matched_book:
                    self._update_book_links(matched_book, row_data)
                    results["matched"] += 1
                    results["details"].append(
                        {"filename": row_data.get("文件名", ""), "book_title": matched_book.title, "status": "matched"}
                    )
                    logger.info(f"匹配成功: {row_data.get('文件名')} -> {matched_book.title}")
                else:
                    results["unmatched"] += 1
                    results["details"].append(
                        {"filename": row_data.get("文件名", ""), "book_title": None, "status": "unmatched"}
                    )
                    logger.warning(f"未匹配: {row_data.get('文件名')}")

            logger.info(f"导入完成: 匹配 {results['matched']}/{results['total']}")
            return results

        except Exception as e:
            logger.error(f"导入Excel失败: {excel_path}, 错误: {e}")
            return {"error": str(e)}

    def _match_book(self, row_data: Dict) -> Optional:
        """匹配书籍"""
        filename = row_data.get("文件名", "")

        if not filename:
            return None

        all_books = self.db.get_all_books()

        best_match = None
        best_score = 0

        for book in all_books:
            if book.status not in ["completed"]:
                continue

            score = match_filename_to_book(filename, book.title or "")
            if score > best_score and score >= 0.5:
                best_score = score
                best_match = book

        return best_match

    def _update_book_links(self, book, row_data: Dict):
        """更新书籍的链接"""
        share_url = row_data.get("分享链接", "")
        extract_code = row_data.get("提取码", "")

        if not share_url:
            return

        sku_includes = self._get_sku_includes(share_url, row_data)

        pan_link_sku1 = sku_includes.get("sku1")
        pan_link_sku2 = sku_includes.get("sku2")

        self.db.update_book_output(
            book.id, pan_link_sku1=pan_link_sku1, pan_link_sku2=pan_link_sku2, pan_code=extract_code
        )

        logger.info(f"更新链接成功: {book.title}")

    def _get_sku_includes(self, share_url: str, row_data: Dict) -> Dict:
        """获取SKU对应的链接"""
        return {"sku1": share_url, "sku2": share_url}

    def generate_excel_template(self, output_path: str = None) -> str:
        """生成导入模板"""
        if not output_path:
            output_path = self.output_dir / "links_template.xlsx"

        ensure_dir(str(Path(output_path).parent))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "分享链接"

        headers = ["文件名", "分享链接", "提取码", "备注"]
        for col_idx, header in enumerate(headers, 1):
            ws.cell(1, col_idx, header)

        wb.save(output_path)
        logger.info(f"模板已生成: {output_path}")

        return str(output_path)


def import_links(excel_path: str) -> Dict:
    """导入链接"""
    importer = LinkImporter()
    return importer.import_from_excel(excel_path)


if __name__ == "__main__":
    print("Excel链接导入器测试")
